import os
import pytest
from docx import Document as DocxDocument
from app.infrastructure.persistence.pdf_processor import PDFProcessor

def test_extract_docx_text(tmp_path):
    # 1. Create a dummy docx file
    doc_path = os.path.join(tmp_path, "test.docx")
    doc = DocxDocument()
    doc.add_paragraph("Hello Zyrabit!")
    doc.add_paragraph("This is a second paragraph.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Col1"
    table.cell(0, 1).text = "Col2"
    table.cell(1, 0).text = "Val1"
    table.cell(1, 1).text = "Val2"
    
    doc.save(doc_path)
    
    # 2. Extract text using PDFProcessor
    docs = PDFProcessor.to_markdown_documents(doc_path)
    
    assert len(docs) == 1
    content = docs[0].page_content
    assert "Hello Zyrabit!" in content
    assert "This is a second paragraph." in content
    assert "Col1 | Col2" in content
    assert "Val1 | Val2" in content
    assert docs[0].metadata["source"] == doc_path
