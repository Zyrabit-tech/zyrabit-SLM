"""Local parsers that produce evidence units with human-readable locators."""
from __future__ import annotations

import re
import subprocess
import tempfile
import unicodedata
import uuid
from pathlib import Path
from typing import Iterable, Sequence

from app.node.domain import EvidenceUnit, Source


class TesseractOcrAdapter:
    def extract(self, path: str, page: int) -> str:
        """OCR one PDF page locally. Requires explicitly installed `tesseract`."""
        try:
            import fitz
            from PIL import Image
        except ImportError as exc:
            raise RuntimeError("OCR dependencies are not installed") from exc
        document = fitz.open(path)
        pix = document.load_page(page).get_pixmap(matrix=fitz.Matrix(2, 2))
        with tempfile.NamedTemporaryFile(suffix=".png") as output:
            Image.frombytes("RGB" if pix.n == 3 else "RGBA", [pix.width, pix.height], pix.samples).save(output.name)
            result = subprocess.run(["tesseract", output.name, "stdout", "-l", "eng+spa"], capture_output=True, text=True, check=False)
        if result.returncode:
            raise RuntimeError(result.stderr.strip() or "Tesseract failed")
        return result.stdout.strip()


class LocalDocumentParser:
    def __init__(self, ocr: TesseractOcrAdapter | None = None):
        self.ocr = ocr

    def parse(self, source: Source, document_id: str) -> Sequence[EvidenceUnit]:
        extension = Path(source.filename).suffix.lower()
        if extension == ".pdf": return list(self._pdf(source, document_id))
        if extension in {".txt", ".md"}: return list(self._text(source, document_id))
        if extension == ".docx": return list(self._docx(source, document_id))
        if extension == ".csv": return list(self._csv(source, document_id))
        if extension == ".xlsx": return list(self._xlsx(source, document_id))
        if extension == ".pptx": return list(self._pptx(source, document_id))
        raise ValueError(f"Unsupported document type: {extension}")

    def _unit(self, document_id: str, text: str, ordinal: int, locator: dict, source: Source) -> EvidenceUnit:
        return EvidenceUnit(id=str(uuid.uuid4()), document_id=document_id, content=self._clean_text(text), ordinal=ordinal,
            locator=locator, metadata={"filename": source.filename, "sha256": source.sha256, "format": Path(source.filename).suffix.lower().lstrip(".")})

    @staticmethod
    def _clean_text(text: str) -> str:
        """Preserve document structure while removing extractor artifacts."""
        text = unicodedata.normalize("NFKC", text)
        text = text.replace("\u00a0", " ").replace("\u00ad", "")
        text = re.sub(r"[\u200b-\u200d\ufeff]", "", text)
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    def _pdf(self, source: Source, document_id: str) -> Iterable[EvidenceUnit]:
        import fitz
        document = fitz.open(source.stored_path)
        for number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            ocr_used = False
            if not text and self.ocr:
                text = self.ocr.extract(source.stored_path, number - 1); ocr_used = True
            if text:
                yield self._unit(document_id, text, number, {"page": number, "ocr": ocr_used}, source)
        if document.page_count and not any(page.get_text("text").strip() for page in document) and not self.ocr:
            raise ValueError("PDF has no extractable text and OCR is disabled")

    def _text(self, source: Source, document_id: str) -> Iterable[EvidenceUnit]:
        text = Path(source.stored_path).read_text(encoding="utf-8", errors="replace").strip()
        if text: yield self._unit(document_id, text, 1, {"section": "document"}, source)

    def _docx(self, source: Source, document_id: str) -> Iterable[EvidenceUnit]:
        from docx import Document
        document = Document(source.stored_path); ordinal = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                ordinal += 1; yield self._unit(document_id, text, ordinal, {"paragraph": ordinal}, source)
        for table_number, table in enumerate(document.tables, start=1):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            text = "\n".join(row for row in rows if row.strip())
            if text:
                ordinal += 1; yield self._unit(document_id, text, ordinal, {"table": table_number}, source)

    def _csv(self, source: Source, document_id: str) -> Iterable[EvidenceUnit]:
        import csv
        with open(source.stored_path, newline="", encoding="utf-8-sig", errors="replace") as handle:
            reader = csv.reader(handle); rows = list(reader)
        if rows:
            text = "\n".join(" | ".join(row) for row in rows)
            yield self._unit(document_id, text, 1, {"sheet": "CSV", "range": f"A1:{len(rows)}"}, source)

    def _xlsx(self, source: Source, document_id: str) -> Iterable[EvidenceUnit]:
        from openpyxl import load_workbook
        workbook = load_workbook(source.stored_path, read_only=True, data_only=True); ordinal = 0
        for worksheet in workbook.worksheets:
            rows = [["" if cell is None else str(cell) for cell in row] for row in worksheet.iter_rows(values_only=True)]
            text = "\n".join(" | ".join(row) for row in rows if any(row))
            if text:
                ordinal += 1; yield self._unit(document_id, text, ordinal, {"sheet": worksheet.title, "range": worksheet.calculate_dimension()}, source)

    def _pptx(self, source: Source, document_id: str) -> Iterable[EvidenceUnit]:
        from pptx import Presentation
        presentation = Presentation(source.stored_path)
        for number, slide in enumerate(presentation.slides, start=1):
            values = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            notes = getattr(slide, "notes_slide", None)
            if notes:
                values.extend(shape.text.strip() for shape in notes.shapes if hasattr(shape, "text") and shape.text.strip())
            text = "\n".join(values)
            if text: yield self._unit(document_id, text, number, {"slide": number}, source)
